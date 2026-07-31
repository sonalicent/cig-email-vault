"""Shared pytest fixtures and synthetic attachment builders.

Attachment binaries (xlsx/docx/pptx) are generated in-memory with their own libraries so
no real files or PII are committed. The PDF builder emits a minimal valid PDF with a
correct xref table so pdfminer/pdfplumber can read it.
"""

from __future__ import annotations

import io
from pathlib import Path

import pytest

FIXTURES_DIR = Path(__file__).parent / "fixtures"


def read_fixture(name: str) -> bytes:
    return (FIXTURES_DIR / name).read_bytes()


@pytest.fixture
def fixture_bytes():
    return read_fixture


def make_txt(text: str = "hello attachment") -> bytes:
    return text.encode("utf-8")


def make_xlsx(rows: list[list[object]] | None = None) -> bytes:
    import openpyxl

    rows = rows or [["Name", "Value"], ["Alpha", 1], ["Beta", 2]]
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def make_docx(paragraphs: list[str] | None = None) -> bytes:
    from docx import Document

    paragraphs = paragraphs or ["First paragraph.", "Second paragraph."]
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_pptx(slide_texts: list[str] | None = None) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    slide_texts = slide_texts or ["Slide one text", "Slide two text"]
    presentation = Presentation()
    blank_layout = presentation.slide_layouts[6]
    for text in slide_texts:
        slide = presentation.slides.add_slide(blank_layout)
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
        box.text_frame.text = text
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def make_pdf(text: str = "Hello PDF") -> bytes:
    """Build a minimal single-page PDF with a correct cross-reference table."""
    stream = f"BT /F1 24 Tf 72 700 Td ({text}) Tj ET".encode("latin-1")
    objects = [
        b"<</Type/Catalog/Pages 2 0 R>>",
        b"<</Type/Pages/Kids[3 0 R]/Count 1>>",
        b"<</Type/Page/Parent 2 0 R/MediaBox[0 0 612 792]"
        b"/Contents 4 0 R/Resources<</Font<</F1 5 0 R>>>>>>",
        b"<</Length %d>>\nstream\n" % len(stream) + stream + b"\nendstream\n",
        b"<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>",
    ]

    pdf = b"%PDF-1.4\n"
    offsets: list[int] = []
    for index, obj in enumerate(objects, start=1):
        offsets.append(len(pdf))
        pdf += b"%d 0 obj\n" % index + obj + b"\nendobj\n"

    xref_pos = len(pdf)
    size = len(objects) + 1
    pdf += b"xref\n0 %d\n" % size
    pdf += b"0000000000 65535 f \n"
    for offset in offsets:
        pdf += b"%010d 00000 n \n" % offset
    pdf += b"trailer<</Size %d/Root 1 0 R>>\nstartxref\n%d\n%%%%EOF" % (size, xref_pos)
    return pdf
