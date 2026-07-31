"""Word (.docx) paragraph + table extraction via python-docx."""

from __future__ import annotations

import io

from docx import Document

from ._util import ExtractionResult, rows_to_markdown_table


def extract(data: bytes) -> ExtractionResult:
    document = Document(io.BytesIO(data))
    parts: list[str] = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)

    for table in document.tables:
        rows = [[cell.text for cell in row.cells] for row in table.rows]
        markdown_table = rows_to_markdown_table(rows)
        if markdown_table:
            parts.append(markdown_table)

    return ExtractionResult(markdown="\n\n".join(parts).strip())
