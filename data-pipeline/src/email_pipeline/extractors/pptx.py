"""PowerPoint (.pptx) slide + speaker-notes extraction via python-pptx."""

from __future__ import annotations

import io

from pptx import Presentation

from ._util import ExtractionResult


def extract(data: bytes) -> ExtractionResult:
    presentation = Presentation(io.BytesIO(data))
    parts: list[str] = []

    for index, slide in enumerate(presentation.slides, start=1):
        parts.append(f"## Slide {index}")

        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    parts.append(text)
            if shape.has_table:
                parts.append(_table_to_markdown(shape.table))

        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(f"**Notes:** {notes}")

    return ExtractionResult(markdown="\n\n".join(parts).strip())


def _table_to_markdown(table) -> str:
    from ._util import rows_to_markdown_table

    rows = [[cell.text for cell in row.cells] for row in table.rows]
    return rows_to_markdown_table(rows)
